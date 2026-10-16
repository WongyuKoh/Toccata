"""Build the final purchase sheet: one sheet, items grouped by store, grand total at the bottom.

usage: python3 build_order_xlsx.py [state_main.json] [out.xlsx]      (EXPECT=731485 to assert the total)
Source: site/index.html (embedded data of https://claude.ai/artifact/9CJ3HwhbhoodCAui94g28n)
      + the db doc state/main of that artifact (items the user turned off = owned / not bought, edited quantities).
Layout follows the user's template hardware/mechanical/구입항목_정리표.xlsx (2026-10-06):
품목 · 단가 · 갯수 · 금액(=단가×갯수) · 사이트 링크, header fill #0B6584, money format #,##0"원".
After building, run the xlsx skill's recalc.py on the output.
"""
import datetime
import json
import os
import re
import sys
from urllib.parse import unquote, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.normpath(os.path.join(HERE, "..", "..", ".."))
STATE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "state_main.json")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(REPO, "최종본", "01_구매목록", "Toccata_구입항목_정리표_최종.xlsx")
EXPECT = int(os.environ["EXPECT"]) if os.environ.get("EXPECT") else None   # e.g. EXPECT=731485 for db state/main v43 (2026-10-02)

page = open(os.path.join(HERE, "site", "index.html"), encoding="utf-8").read()
D = json.loads(re.search(r'<script type="application/json" id="data">(.*?)</script>', page, re.S).group(1))
S = json.load(open(STATE, encoding="utf-8"))

# store display names / where to order, store notes and per-item notes (shared with the final-BOM page export)
NOTES = json.load(open(os.path.join(HERE, "order_notes.json"), encoding="utf-8"))
SITE = {k: tuple(v) for k, v in NOTES["site"].items()}
STORE_NOTE, NOTE = NOTES["store_note"], NOTES["note"]


def unit_at(it, q):
    u = it["unit"]
    for mq, up in it.get("tiers") or []:
        if q >= mq:
            u = up
    return u


stores = {}
for it in D["items"]:
    if not S["on"].get(it["id"], it["on"]):
        continue
    q = S["qty"].get(it["id"], it["qty"])
    u = unit_at(it, q)
    if u * q == 0:                      # FluidSynth / Salamander: free downloads, nothing to order
        continue
    k = it.get("sk") or it["store"]
    st = stores.setdefault(k, {"key": k, "rows": [], "sub": 0})
    st["rows"].append((it, q, u))
    st["sub"] += u * q
for st in stores.values():
    sh = D["ships"].get(st["key"]) or {"fee": 0, "free_over": None, "verified": True}
    st["ship"] = sh
    fo = sh.get("free_over")
    carry = sh.get("carry", 0)          # other project's items in the same cart (aircon list)
    st["fee_now"] = 0 if (sh["fee"] == 0 or sh.get("shared") or (fo is not None and st["sub"] + carry >= fo)) else sh["fee"]
n_off = sum(1 for it in D["items"] if not S["on"].get(it["id"], it["on"]))
parts = sum(s["sub"] for s in stores.values())
ship = sum(s["fee_now"] for s in stores.values())
assert EXPECT is None or parts + ship == EXPECT, (parts, ship, parts + ship)
order = sorted(stores.values(), key=lambda s: -(s["sub"] + s["fee_now"]))

# ---- styles (from the user's template) ----
FONT = "Apple SD Gothic Neo"
TEAL, BAND, LINE = "0B6584", "E7EEF2", "B7C9D3"
f_base = Font(name=FONT, size=11)
f_in = Font(name=FONT, size=11, color="0000FF")                 # editable inputs: 단가 · 갯수
f_head = Font(name=FONT, size=11, bold=True, color="FFFFFF")
f_store = Font(name=FONT, size=12, bold=True, color="0B3B4F")
f_small = Font(name=FONT, size=10, color="5A6B75")
f_ship = Font(name=FONT, size=10.5, color="5A6B75")
f_sub = Font(name=FONT, size=11, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color="0B3B4F")
f_link = Font(name=FONT, size=10.5, color="0563C1", underline="single")
f_total = Font(name=FONT, size=14, bold=True, color="FFFFFF")
fill_head = PatternFill("solid", fgColor=TEAL)
fill_band = PatternFill("solid", fgColor=BAND)
thin = Side(style="thin", color=LINE)
b_bottom = Border(bottom=thin)
b_top = Border(top=Side(style="thin", color="5A6B75"))
WON = '#,##0"원"'
WIDTHS = {"A": 46, "B": 13, "C": 8, "D": 15, "E": 48, "F": 46}

wb = Workbook()
ws = wb.active
ws.title = "구입항목"
for c, w in WIDTHS.items():
    ws.column_dimensions[c].width = w
ws.sheet_view.showGridLines = False


def put(r, c, v, font=f_base, fmt=None, al=None, fill=None, border=None):
    cell = ws.cell(r, c, v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if al:
        cell.alignment = al
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    return cell


L = Alignment(horizontal="left", vertical="center", wrap_text=True)
R = Alignment(horizontal="right", vertical="center")
C = Alignment(horizontal="center", vertical="center")

put(1, 1, "Toccata 구입 항목 정리표 (최종)", f_title)
ws.row_dimensions[1].height = 28
put(2, 1, f"{datetime.date.today().isoformat()} 정리 · 최종 재료 페이지에서 고른 그대로(가진 부품·공구 빼고, 예비 최소, 판매처 묶기·무료배송 맞추기 반영) · "
          "단가는 VAT 포함 확인값이라 주문 전에 다시 확인 · 파란 글씨(단가·갯수)를 고치면 금액·배송비·합계가 다시 계산됩니다", f_small,
    al=Alignment(wrap_text=True, vertical="top"))
ws.merge_cells("A2:F2")
ws.row_dimensions[2].height = 30

HR = 4
for i, h in enumerate(["품목", "단가", "갯수", "금액", "사이트 링크", "비고"], 1):
    put(HR, i, h, f_head, al=C, fill=fill_head)
ws.row_dimensions[HR].height = 25
ws.freeze_panes = ws.cell(HR + 1, 1)

r = HR + 1
item_ranges, ship_cells = [], []
for n, st in enumerate(order, 1):
    name, site = SITE.get(st["key"]) or (st["key"], urlparse(st["rows"][0][0].get("url") or "").netloc.replace("www.", ""))
    sh = st["ship"]
    fo = sh.get("free_over")
    if sh["fee"] == 0:
        rule = "무료배송"
    elif fo:
        rule = f"배송비 {sh['fee']:,}원 · {fo:,}원 이상 무료"
    else:
        rule = f"배송비 {sh['fee']:,}원 (무료 조건 없음)"
    if sh.get("shared"):
        rule = f"배송비 {sh['fee']:,}원은 에어컨 목록에 있음 (같은 장바구니)"
    elif sh.get("carry"):
        rule += f" · 에어컨 재료 {sh['carry']:,}원과 같은 장바구니"
    if not sh.get("verified", True):
        rule += " · 추정"
    r += 1                                               # blank spacer row between stores
    head = f"{n}. {name}" + (f"  ({site})" if site else "")
    put(r, 1, head, f_store, fill=fill_band, al=Alignment(vertical="center"))
    put(r, 2, None, fill=fill_band)
    put(r, 3, None, fill=fill_band)
    put(r, 4, None, fill=fill_band)
    put(r, 5, f"{len(st['rows'])}품목 · {rule}", f_small, fill=fill_band, al=Alignment(vertical="center"))
    put(r, 6, STORE_NOTE.get(st["key"], ""), f_small, fill=fill_band, al=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[r].height = 24
    r += 1
    first = r
    for it, q, u in sorted(st["rows"], key=lambda x: -x[1] * x[2]):
        url = it.get("url", "")
        put(r, 1, it.get("t") or it["n"], al=L, border=b_bottom)
        put(r, 2, u, f_in, WON, R, border=b_bottom)
        put(r, 3, q, f_in, "#,##0", C, border=b_bottom)
        put(r, 4, f"=B{r}*C{r}", f_base, WON, R, border=b_bottom)
        if url.startswith("http"):
            p = urlparse(url)
            shown = unquote(p.netloc.replace("www.", "") + p.path + (("?" + p.query) if p.query else ""))
            if len(shown) > 46:
                shown = shown[:45] + "…"
            cell = put(r, 5, shown, f_link, al=Alignment(vertical="center"), border=b_bottom)
            cell.hyperlink = url
        else:
            put(r, 5, "", border=b_bottom)
        put(r, 6, NOTE.get(it["id"], ""), f_small, al=L, border=b_bottom)
        r += 1
    last = r - 1
    item_ranges.append(f"D{first}:D{last}")
    # shipping row: live formula from the store's own subtotal
    put(r, 1, "배송비", f_ship, al=Alignment(horizontal="right", vertical="center"))
    if sh["fee"] == 0 or sh.get("shared"):
        put(r, 4, 0, f_ship, WON, R)
    elif fo:
        carry = sh.get("carry", 0)
        put(r, 4, f"=IF(SUM(D{first}:D{last})+{carry}>={fo},0,{sh['fee']})" if carry else f"=IF(SUM(D{first}:D{last})>={fo},0,{sh['fee']})", f_ship, WON, R)
    else:
        put(r, 4, sh["fee"], f_ship, WON, R)
    ship_cells.append(f"D{r}")
    r += 1
    put(r, 1, f"{name.split(' · ')[0]} 소계 (물건값 + 배송비)", f_sub, al=Alignment(horizontal="right", vertical="center"), border=b_top)
    for c in (2, 3, 5, 6):
        put(r, c, None, border=b_top)
    put(r, 4, f"=SUM(D{first}:D{last})+D{r - 1}", f_sub, WON, R, border=b_top)
    r += 1

# ---- grand total at the bottom ----
r += 1
TOT = r
n_items = sum(len(s["rows"]) for s in order)
put(r, 1, f"물건값 합계 ({n_items}품목)", f_sub, al=Alignment(horizontal="right", vertical="center"), fill=fill_band)
for c in (2, 3, 5, 6):
    put(r, c, None, fill=fill_band)
put(r, 4, "=SUM(" + ",".join(item_ranges) + ")", f_sub, WON, R, fill=fill_band)
r += 1
put(r, 1, f"배송비 합계 (판매처 {len(order)}곳)", f_sub, al=Alignment(horizontal="right", vertical="center"), fill=fill_band)
for c in (2, 3, 5, 6):
    put(r, c, None, fill=fill_band)
put(r, 4, "=" + "+".join(ship_cells), f_sub, WON, R, fill=fill_band)
r += 1
put(r, 1, "총합 (최종 가격)", f_total, al=Alignment(horizontal="right", vertical="center"), fill=fill_head)
for c in (2, 3, 5, 6):
    put(r, c, None, fill=fill_head)
put(r, 4, f"=D{TOT}+D{TOT + 1}", f_total, WON, R, fill=fill_head)
put(r, 5, "VAT 포함 · 배송비 포함", Font(name=FONT, size=10.5, color="FFFFFF"), fill=fill_head, al=Alignment(vertical="center"))
ws.row_dimensions[r].height = 30
r += 2
for line in [
    f"· 이 목록에 없는 것: 최종 재료 페이지에서 끈 품목 {n_off}가지(이미 가진 부품·공구, 설계에서 빠진 것)와 무료 소프트웨어(FluidSynth, Salamander 피아노 음원 — 다운로드).",
    f"· 출처: 최종 재료 페이지 https://claude.ai/artifact/9CJ3HwhbhoodCAui94g28n (품목 사진·규격·쓰는 위치) + 사용자가 고른 상태(db state/main, 저장 {str(S.get('updatedAt', ''))[:10]}).",
    "· 배송비 줄: 무료 기준이 있는 판매처는 '그 판매처 물건값 ≥ 기준이면 0원' 수식입니다. 기준 금액은 VAT 포함으로 바꿔 적었습니다(예: 5만 원 VAT 별도 → 55,000원).",
    "· 홀센서(100개)와 미스미 스프링(100개)은 구간 단가라, 갯수를 100개 밑으로 줄이면 단가가 올라갑니다(홀센서 1,023원, 스프링 365원) — 단가도 함께 고치세요.",
    "· 생성기: hardware/bom/final-bom/build_order_xlsx.py",
]:
    put(r, 1, line, f_small)
    r += 1

ws.print_title_rows = f"{HR}:{HR}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.4

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT, "items", n_items, "stores", len(order), "parts", parts, "ship", ship, "total", parts + ship)
