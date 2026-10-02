"""Toccata v4 재료 목록 + 절감 선택기 데이터.

기준: 사용자 구매 목록(docs/report/toccata-purchase-list.xlsx, v3.2 + 22개 절감 선택)에
v4 W1+(r3) 건반 액션 변경을 반영한 목록. 대안은 옛 절감안(도면집 v10 선택기)과
v4 부품 조사(costopt/)에서 온다. 2026-09-28 회로 검토(hardware/pcb README)의 줄 변경과 USB 케이블 조사(../cable/result.json)도 여기서 넣는다.
출력: data_old.json (카드·옵션·배송 규칙·사진 키), imgs_old.json
"""
import base64, glob, html, io, json, os, re
from collections import OrderedDict
from openpyxl import load_workbook
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
V4 = os.path.dirname(HERE)
REPO = '/Users/kwg/Desktop/mydrive/project/Toccata'
SEL = '/private/tmp/claude-501/-Users-kwg-Desktop-mydrive-project-Toccata/4d582017-bb60-49f5-9df1-1f2d9e6b1771/scratchpad/sel'
IMG220 = '/private/tmp/claude-501/-Users-kwg-Desktop-mydrive-project-Toccata/4115e6ae-9485-4e49-9b1a-8afa15589664/scratchpad/v3/cost/img220'
DOC = os.path.join(REPO, 'docs/report/toccata-build-v3.html')
XLSX = os.path.join(REPO, 'docs/report/toccata-purchase-list.xlsx')

# The old picker data (session 4d582017 scratchpad) can disappear when that temp folder is cleaned. Then run in
# CACHED mode: take the old alternatives and their photos from the last good output (data_old.cache.json /
# imgs_old.cache.json) and rebuild everything else (rows, current options, patches) from the sources as usual.
CACHED = not os.path.exists(os.path.join(SEL, 'opt_debug.json'))
if CACHED:
    OPT, DRAFT = {}, {}
    _PRIOR = json.load(open(os.path.join(HERE, 'data_old.cache.json')))
    _PRIOR_IMGS = json.load(open(os.path.join(HERE, 'imgs_old.cache.json')))
else:
    OPT = json.load(open(os.path.join(SEL, 'opt_debug.json')))['OPT']
    DRAFT = {}
    for f in sorted(glob.glob(os.path.join(SEL, 'drafts', 'out_*.json'))):
        for d in json.load(open(f)):
            DRAFT[d['id']] = d
DOCS = open(DOC, encoding='utf-8').read()

# ---------------------------------------------------------------- images
IMGS = {}
if CACHED:
    IMGS.update(_PRIOR_IMGS)


def thumb(raw):
    im = Image.open(io.BytesIO(raw)).convert('RGB')
    im.thumbnail((160, 160))
    bg = Image.new('RGB', (160, 160), (255, 255, 255))
    bg.paste(im, ((160 - im.width) // 2, (160 - im.height) // 2))
    out = io.BytesIO()
    bg.save(out, 'JPEG', quality=72, optimize=True, progressive=True)
    return base64.b64encode(out.getvalue()).decode()


def img_from_bytes(key, raw):
    if not raw:
        return None
    try:
        IMGS[key] = thumb(raw)
        return key
    except Exception:
        return None


_IMG_CACHE_V4 = json.load(open(os.path.join(HERE, 'imgs_v4.cache.json'))) if os.path.exists(os.path.join(HERE, 'imgs_v4.cache.json')) else {}


def img_file(key, path):
    if path and os.path.exists(path):
        return img_from_bytes(key, open(path, 'rb').read())
    if key in _IMG_CACHE_V4:  # 10/2: source photos in scratch v4/costopt/img were purged by the OS tmp cleaner — reuse the last built thumbnail
        IMGS[key] = _IMG_CACHE_V4[key]
        return key
    return None


def img_line(lid):
    i = DOCS.find(f'<article class="pc-card" id="{lid}"')
    if i < 0:
        return None
    c = DOCS[i:DOCS.find('</article>', i)]
    m = re.search(r'<img src="data:image/[a-z]+;base64,([A-Za-z0-9+/=]+)"', c)
    return img_from_bytes('L-' + lid, base64.b64decode(m.group(1))) if m else None


def img_opt(oid, idx):
    it = OPT[oid]['add'][idx]
    return img_file(f'N-{it["src"]}-{it["idx"]}', os.path.join(IMG220, f't_{it["src"]}_{it["idx"]}.jpg'))


# ---------------------------------------------------------------- shipping groups
# key -> fee, bucket, threshold (VAT-excluded when vat=True), label
SHIP = OrderedDict([
    ('아이씨뱅큐', dict(fee=2700, b='base', thr=50000, vat=True, label='아이씨뱅큐 (공급가 5만원 미만)')),
    ('엘레파츠', dict(fee=3000, b='consumable', thr=60000, vat=True, label='엘레파츠 (공급가 6만원 미만)')),
    ('Bambu Lab', dict(fee=3000, b='consumable', thr=70000, vat=False, label='Bambu Lab (7만원 미만)')),
    ('세운상가새한음향', dict(fee=4000, b='base', label='11번가 세운상가새한음향')),
    ('두유레디', dict(fee=3000, b='base', label='11번가 두유레디')),
    ('상우아트', dict(fee=3000, b='base', label='11번가 상우아트')),
    ('오공 본드 판매처', dict(fee=2500, b='consumable', label='11번가 오공 본드 판매처')),
    ('굿나잇몰', dict(fee=3500, b='base', label='11번가 굿나잇몰')),
    ('메카솔루션', dict(fee=3000, b='base', label='메카솔루션')),
    ('파츠파츠', dict(fee=3000, b='base', label='파츠파츠 직영몰')),
    ('다나와 최저가 판매처', dict(fee=3000, b='base', label='다나와 최저가 판매처')),
    ('대건상사', dict(fee=2500, b='base', label='대건상사')),
    ('알씨뱅크', dict(fee=3500, b='base', label='알씨뱅크')),
    ('철물박사', dict(fee=3000, b='base', label='철물박사')),
    ('스튜디오분트', dict(fee=3000, b='base', label='스튜디오분트')),
    ('피아노모아', dict(fee=3000, b='base', label='피아노모아')),
    ('PINE STORE', dict(fee=16200, b='tool', label='PINE STORE (해외)')),
    ('DigiKey', dict(fee=20000, b='consumable', thr=60000, vat=False, label='DigiKey Korea (6만원 미만)')),
])


def ship_key(seller, section=''):
    s = (seller or '') + ' ' + (section or '')
    for k in SHIP:
        if k in s:
            return k
    if '파츠파츠' in s:
        return '파츠파츠'
    return None


BK = {'기본': 'base', '소모품': 'consumable', '공구': 'tool'}

# ---------------------------------------------------------------- current list rows (v3.2 purchase list)
wb = load_workbook(XLSX)
ws = wb['구매목록']
ROWS = []
sec = None
for r in range(6, 160):
    a = ws.cell(r, 1).value
    if isinstance(a, str) and a.startswith('■'):
        sec = a[2:].strip()
        continue
    name = ws.cell(r, 2).value
    if name is None or ws.cell(r, 5).value == '배송비' or str(name).startswith('배송비'):
        continue
    code = ws.cell(r, 11).value
    if code is None:
        continue
    ROWS.append(dict(row=r, section=sec, name=name, unit=ws.cell(r, 3).value, qty=ws.cell(r, 4).value,
                     role=ws.cell(r, 5).value or '', url=ws.cell(r, 7).value or '', b=BK[ws.cell(r, 8).value],
                     seller=ws.cell(r, 10).value or '', code=code, note=ws.cell(r, 12).value or ''))

# ---------------------------------------------------------------- v4 changes to existing rows
REMOVED_V4 = {  # code -> reason
    'L32': '척추 클램프 나사 — v4에서 리프 콤·척추가 없어져 빠짐',
    'L33': '무게 조절 세트스크루 — v4에서 빠짐(무게는 캡스턴·텅스텐 퍼티로)',
    'L38': '세트스크루 받침 황동 원판 — v4에서 빠짐',
}
for r in ROWS:
    if r['code'] == 'T11' and '1.5' in r['name']:
        r['drop'] = '1.5 mm 렌치(세트스크루용) — v4에서 빠짐'
    if r['code'] in REMOVED_V4:
        r['drop'] = REMOVED_V4[r['code']]
    if r['code'] == 'L36':
        r['qty'] = 20
        r['role'] = '고무발 고정 피스 머리 밑 와셔 14 + 예비 6 (v4: 클램프 나사용 32개는 빠짐)'
    if r['code'] == 'L31':
        r['qty'] = 100
        r['role'] = 'v4: 건반 캡스턴 나사(M3 버튼헤드) 너트 88 + 예비 8, 스피커 브래킷 4'
    if r['code'] == 'ME8':
        r['role'] = '센서 자석 Ø5×2·뚜껑 자석을 포켓에 압입한 뒤 한 방울, 건반 봉 끝을 봉 받침에 고정(v4)'
    if r['code'] == 'CT2' and '블랙' in r['name'] and '스풀 포함' in r['name']:
        r['qty'] = 5
        r['role'] = 'v4: 프레임·레버 캐리어·흑건·커버 등. v4 출력물이 약 6.5 kg(서포트 포함)으로 늘어 블랙 2 kg 추가(3 → 5)'
    if r['code'] == 'L39':
        r['role'] = '건반 앞 멈춤 펠트 3T(백·흑)와 키퍼 훅 밑 펠트. v4에서도 그대로 쓴다'

# ---------------------------------------------------------------- circuit review (2026-09-28)
# hardware/pcb/README.md "구매 목록(v4)에 넣거나 고칠 줄". New rows L67·L68 (codes given here; bom.csv says "신규").
# Prices checked on the 아이씨뱅큐 pages 2026-09-28 (VAT 포함, 공급가 5만원 이상 무료배송, 미만 2,700원).
CIRC = os.path.join(HERE, 'circuit')
ROWS += [
    dict(row=None, section='아이씨뱅큐', name='C3216-10UF(K 16V) — 1206 10 µF 16 V 칩 세라믹 콘덴서 (50개 단위)', unit=7150, qty=1,
         role='센서 기판 C213·C214 ×7 + 끝 부속 EL C404 + ER C412 = 16개(회로 △1·△3). 34개 남음',
         url='https://www.icbanq.com/P001248090', b='base', seller='아이씨뱅큐 (오렌지반도체)', code='L67',
         note='9/28 확인: 6,500원 + VAT = 7,150원, 국내 재고, 1~2일. 제조사·유전체 표기 없음. 주문 제작 상품이라 취소·반품 불가',
         img_path=os.path.join(CIRC, 'img_10uF.jpg')),
    dict(row=None, section='아이씨뱅큐', name='Yageo RC1206JR-071KL — 1206 1 kΩ 5 % 1/4 W 칩 저항', unit=44, qty=10,
         role='끝 부속 EL R401~R403 + ER R411 = 4개(회로 △13, 기판 아랫면). 10개 단위 판매',
         url='https://www.icbanq.com/P011433924', b='base', seller='아이씨뱅큐 (element14 해외 재고)', code='L68',
         note='9/28 확인: 40원 + VAT = 44원, 10개 단위, 해외 재고 4,690개, 4~6일. 해외 재고는 취소·반품 불가',
         img_path=os.path.join(CIRC, 'img_1k.jpg')),
]
L29_ROLE = {'1kΩ': '1 kΩ ×4: R701·R702(앰프 입력), R501·R551(페달). 6개 남음',
            '10kΩ': '회로에서 쓰지 않음(예비). 빼도 된다',
            '100kΩ': '100 kΩ ×5: R301(O1·O7)·R302·R303(제어 기판), R502(페달). 5개 남음',
            '100Ω': '100 Ω ×5 병렬 = 20 Ω: R503(페달, 필수). 5개 남음'}


def swap(r, old, new):
    assert old in r['role'], (r['code'], old)
    r['role'] = r['role'].replace(old, new)


for r in ROWS:
    c, n = r['code'], str(r['name'])
    if c == 'SE7' and '소켓' in n:
        r['drop'] = 'EXT 암 커넥터용 — 회로 검토에서 빠짐(EXT는 L66 JST-XH 리드를 패드에 직접 납땜)'
    elif c == 'SE7':
        r['qty'] = 9
        r['role'] = '4067 먹스 7개 × (16+8)핀 = 168핀 + RP2040-Zero 8개 × 1×9 ×2 = 144핀 → 312핀, 1×40 9줄(360핀)'
    if c == 'SEV2' and 'RP2040' in n:
        swap(r, '캐스털레이션으로 제어 기판에 직접 납땜', '안쪽 구멍 핀 헤더 1×9 ×2(몰드 뺌)로 제어 기판에 납땜, 캡톤 위 약 +1.3 mm')
    if c == 'L29':
        r['role'] = L29_ROLE[re.search(r'저항 (\S+) \(', n).group(1)]
    if c == 'C05':
        r['role'] = '센서 기판 → 제어 기판 리본 16심 전부(OUT 12 + 3V3 ×2 + GND ×2, 약 180 mm) ×7 + EL·ER 리드 가닥'
    if c == 'L66':
        swap(r, '연장선 가운데를 잘라', '200 mm 연장선을 모듈 쪽 150 · 끝 부속 쪽 50으로 잘라')
    if c == 'C04':
        swap(r, '몰드 ≤12.5×7.5×25라 통로(높이 22)에 들어간다', '모듈 쪽 C 플러그 몰드 ≤ 12.5 × 7.6 × 25(폭 초과 금지)')
        r['note'] = ('NA993은 몰드 두께가 공개되지 않았다. 폭 약 11.3·길이 약 19.5는 사진으로 잰 추정값이다. '
                     '단계 0 시험 19(프레임 뒤 조각에 10번 꽂고 빼기)로 확인한다')
    if c == 'L56':
        swap(r, '약 80 Hz', '약 93 Hz')
# ---------------------------------------------------------------- r4.4 check (2026-09-28)
# EXT 헤더 커넥터는 빠졌다: EXT는 L66 JST-XH 리드를 패드에 직접 납땜한다. L35는 r4.4에서도 모듈당 2개(14 + 예비 6)로 충분.
for r in ROWS:
    if r['code'] == 'C10':
        swap(r, 'EXT 헤더 커넥터', 'EXT 리드(L66 JST-XH)를 패드에 납땜한 곳')
    if r['code'] == 'C12':
        swap(r, 'EXT 커넥터 고정', 'EXT 리드(L66) 당김 방지')
    if r['code'] == 'L35':
        swap(r, '대각 2곳만 나사, 나머지 2곳은 출력 위치 핀(v3.2)',
             '모듈마다 대각 2곳만 나사(7 × 2 = 14 + 예비 6), 나머지 2곳은 출력 위치 핀. r4.4도 같다')
# ---------------------------------------------------------------- r4.5 circuit 2nd check (2026-09-29, 회로 세션 2차 요청 12)
for r in ROWS:
    if r['code'] == 'L35':
        r['role'] = ('제어 기판을 프레임의 매단 보스 4개에 고정 — 모듈마다 대각 2곳을 밑에서 나사로 조임(7 × 2 = 14 + 예비 6), '
                     '나머지 2곳은 보스의 위치 핀. r4.5: 바닥 구멍으로 기판을 밑에서 넣음. 스텐 유두(ISO 7380, 머리 Ø5.7 × 1.65)')
    if r['code'] == 'SE4':
        swap(r, '센서마다 1개(3V3–GND 디커플링, 센서 기판 아랫면) + PED 보드',
             '센서마다 1개(3V3–GND 디커플링, 센서 기판 아랫면) + PED 보드 + 제어 기판 C301 ×7 = 97개 사용, 예비 3')
    if r['code'] == 'L20':
        swap(r, '건반마다 1개 + 예비 12', '건반 88 + 페달 1 = 89개 사용, 예비 11')
    if r['code'] == 'L21':
        swap(r, '88 + 예비 12', '건반 88 + 페달 1 = 89개 사용, 예비 11')
# ---------------------------------------------------------------- R30 / CAD L1 (2026-09-30): lower rear bar, slanted speaker baffle
for r in ROWS:
    if r['code'] == 'L10':
        r['role'] = ('스피커 파트 2개의 상자(R30·CAD L1: 외형 213×195×129, 앞판 위쪽 30° 경사, 안 부피 2.95 L)를 한 장에서 재단한다. '
                     '파트마다 옆판 195×129 ×2(앞 윗모서리 30° 한 번 자름), 뒤판 190×129, 윗판 190×109.7, 아랫판 190×183.5, 앞 아래 띠 190×12.6 = 12조각. '
                     '경사 앞판은 PETG 출력. 12조각이 400×1200 한 장의 약 710 mm 길이에 들어간다(재단 톱날 3 mm 포함). 치수 출처: hardware/mechanical/cad/manifest.json')
    if r['code'] == 'L63':
        r['role'] = ('가운데 유닛(R30·CAD L1: 822×195×83, x200~1022) 상자를 한 장에서 재단한다. 바닥 356×195·266×195, 뚜껑 356×195·266×195, 앞판 822×60(아래 홈 x583~733 높이 17.5), '
                     '뒤판 356×60·266×60, 끝판 172×60 ×2, 칸막이 172×60, 낮은 칸막이 172×40 = 11조각(CU 칸 바닥·뚜껑·I/O 판은 출력). 600×1200 한 장에 들어간다')
    if r['code'] == 'L14':
        r['role'] = ('스피커 파트 2개(각 4개)와 가운데 유닛(6개)의 고무발 14개. R30·CAD L1: 뒷바를 5 mm만 띄운다(22 mm 띄움과 4 mm 출력 받침은 없어짐; 케이블은 모듈 뒤 y212~252 열린 통로로). '
                     '본체와 스피커 진동을 떼어 놓는다(D15)')
    if r['code'] == 'L07':
        swap(r, '스피커 파트(오꾸메 밀폐 상자, 내부 5.05 L)의 출력 앞판에 끼워 연주자 쪽을 향하게 한다',
             '스피커 파트(오꾸메 밀폐 상자, 안 부피 2.95 L — R30·CAD L1)의 30° 경사 출력 앞판에 달아 앉은 연주자 귀를 향하게 한다')
    if r['code'] == 'SP6':
        swap(r, '스피커 상자 2개(5.05 L ×2)', '스피커 상자 2개(2.95 L ×2)')
    if r['code'] == 'L49' and '25mm' in str(r['name']):
        r['drop'] = 'R30·CAD L1: 4 mm 출력 받침이 없어져 발 5 + 합판 11.5 = 16.5에 25 mm를 쓰면 끝이 상자 안으로 약 8 mm 뚫고 나온다 → 고무발도 13 mm로'
    if r['code'] == 'L49' and '13mm' in str(r['name']):
        r['qty'] += 14
        r['role'] = (r['role'] + ' + 고무발 14개 고정(R30·CAD L1: 발 머리 자리 Ø9×2 가정 → 발 3 + 합판 10 물림, 끝이 합판 안면 1.5 mm 안. '
                     '받은 고무발의 머리 자리가 2 mm보다 깊으면 16 mm로)')
# ---------------------------------------------------------------- 9/30 판매처 변경 (사용자 결정, 가공 필요 재료 정리 세션 조사; 가격은 W1 세션이 9/30 페이지에서 다시 확인)
for r in ROWS:
    if r['code'] == 'L42':
        r.update(name='알씨뱅크 M3 진동 방지 댐퍼 (20개, Ø6 × 8.6 mm, 허리 홈 FC 댐퍼형)', unit=4500, qty=1,
                 url='https://www.rcbank.co.kr/shop/goods/goods_view.php?goodsno=67110', seller='알씨뱅크', section='알씨뱅크',
                 note='9/30 판매처 변경(이전 AliExpress WorldBridge 3,040원, 해외 배송). 알씨뱅크는 L66 연장선 주문과 같이 와서 배송비가 더 들지 않는다',
                 img_path='/private/tmp/claude-501/-Users-kwg-Desktop-mydrive-project-Toccata/2f19e142-4ac2-4c90-bd1e-b3b469f97487/scratchpad/v4/costsel/vendor930/grommet_rcbank.jpg')
    if r['code'] == 'SP6':
        r.update(name='[오공] 목공용 수성 접착제 205본드 800 g (나비엠알오 K01424797)', unit=3949, qty=1,
                 url='https://www.navimro.com/g/395062/', seller='나비엠알오', section='나비엠알오',
                 note='9/30 판매처 변경(이전 11번가 2,840원 + 배송 2,500원). 나비엠알오는 PORON·하네나이트 주문에 묶여 배송비가 더 들지 않는다. 800 g 품번 K01424797을 고른다',
                 img_path='/private/tmp/claude-501/-Users-kwg-Desktop-mydrive-project-Toccata/2f19e142-4ac2-4c90-bd1e-b3b469f97487/scratchpad/v4/costsel/vendor930/bond205_navimro.jpg')
# ---------------------------------------------------------------- R31 터치스크린 (2026-09-30): CU 뚜껑 앞 모서리 고정 나사
for r in ROWS:
    if r['code'] == 'L01' and 'GPIO 배선 없음' in r['role']:
        r['role'] = r['role'].replace('GPIO 배선 없음', 'GPIO는 2번(5 V)·6번(GND)만 터치스크린 전원으로 씀 — R31, 회로 개정 C △16')
    if r['code'] == 'L34':
        r['qty'] += 4
        r['role'] = r['role'] + ' + R31: CU 뚜껑(터치스크린 받침) 앞 모서리 2곳을 나사형 받침에 고정 + 예비 2'
# ---------------------------------------------------------------- 10/1 L2 (한 몸 뒷바) + 터치스크린 3판 (사용자 승인 2026-10-01)
# 출처: hardware/mechanical/cad/spec/body_L2.json (bom_changes·joints·plywood_cut_list·centre·implementation_notes, 앞판 40°),
#       hardware/mechanical/cad/README.md ('모델에 넣지 않은 것', '구매 목록에 없는 것', 'L2에서 더 이상 사지 않는 것'),
#       hardware/mechanical/touchscreen/rev3/design/CAD_SPEC_rev3.md (뚜껑 M3×10 4개 → 레일 M3 열압입 인서트, 리본 약 256/300).
# 가격은 10/1 판매 페이지에서 확인(VAT 포함, 엘레파츠는 VAT 별도 값 × 1.1). 이미 목록에 있는 판매처만 써서 배송비 줄이 늘지 않는다
# (엘레파츠는 무료 기준 위, 굿나잇몰·나비엠알오·나사코리아는 기존 주문에 묶임). 새 상품 사진은 내려받지 않았다(사진 없음).
L2 = '10/1 L2'
GN_BOLT = 'https://www.11st.co.kr/products/4236198413'  # 굿나잇몰 렌치볼트 (유두 옵션: M3×5 40 · M3×10 60 · M3×16 90 · M3×20 70 · M4×12 80 · M2.5×6 200)
for r in ROWS:
    c, n = r['code'], str(r['name'])
    if c == 'L63':
        r['drop'] = (L2 + ': 가운데 유닛 판이 스피커 판과 함께 400×1200 한 장(L10)에 모두 들어간다. '
                     '가게가 L10 배치대로 자르지 못하면 이 600×1200을 다시 산다(+22,520원)')
    if c == 'L64':
        r['drop'] = (L2 + ': 도브테일 블록·토글 래치 대신 BRK2 세운 핀 + M4 손잡이볼트(L72) 2개씩으로 스피커와 가운데를 잇는다. '
                     '철물박사 주문이 없어져 배송비 3,000원도 빠진다')
    if c == 'L49' and '16mm' in n:
        r['drop'] = L2 + ': 스피커 유닛은 출력 앞판의 M4 열압입 인서트(L69) + M4×12(L70)로 고정한다(직결피스 16 mm 대신)'
    if c == 'L49' and '13mm' in n:
        r['qty'] = 20
        r['role'] = (L2 + ': 고무발 14개 고정 + 예비 6 (둥근머리 13 mm는 고무발에만 쓴다). 16 mm 이상은 스피커 밀폐 바닥을 뚫으므로 13 mm 이하'
                     '(발 머리 자리 Ø9×2 가정 → 발 3 + 합판 10 물림). L1의 도브테일·CU 트레이·XT30 받침은 없어졌고, '
                     '뒤판 출력물·보조배터리 받침은 접시머리(L77), 허브 선반 받침은 M3×12 태핑(L78), 이음 레일·XT30 집게는 MS 폴리머로 붙인다')
        r['note'] = '같은 상품의 옵션. 굿나잇몰 묶음배송(주문 화면에서 묶음 여부 확인). ' + L2 + ' 수량 54 → 20'
    if c == 'L49' and '19mm' in n:
        r['role'] = 'PETG 그릴 8곳 고정(그릴마다 4개, 떨리지 않게 모두). ' + L2 + ': 19 mm 이하 — 25 mm는 앞판 아래 파일럿 끝 2.15 mm 밑 밀폐면을 뚫는다'
    if c == 'L34':
        r['qty'] = 28
        r['role'] = ('센서 바·센서 기판을 왼쪽 받침 기둥에 고정(모듈 7 + 끝 부속 2 = 9곳 × 2 = 18, PETG 구멍에 직접 탭) + '
                     + L2 + '·터치스크린 3판: 화면 뚜껑(CU-SCREENLID)을 이음 레일의 M3 열압입 인서트(L71) 4곳에 고정(앞 y246 2 · 뒤 y322 2, '
                     '자리파기 Ø6.5 × 5.5, 인서트 물림 4) + 예비 6 = 28. L2에서 CU 트레이 부품 나사 14개는 M3×5(L76)·M3×6(L35)로 바뀌고 '
                     '토글 래치(L64) 8개는 없어졌다')
        r['note'] = L2 + ' 수량 54 → 28 (R31 2판의 "CU 뚜껑 앞 모서리 2곳 + 예비 2"는 3판에서 레일 인서트 4곳으로 바뀜)'
    if c == 'L35':
        r['qty'] = 24
        r['role'] = r['role'] + ' + ' + L2 + ': 앰프 보드(XH-A232)를 선반 받침에 4개(받침 + 선반 관통, PETG 직접 탭) → 24'
    if c == 'L36':
        r['role'] = ('고무발 고정 피스 머리 밑 와셔 14 + ' + L2 + ': BRK2 방진 브래킷의 M3×16(L75) 머리 밑 4 + 예비 2 '
                     '(v4: 클램프 나사용 32개는 빠짐)')
    if c == 'L47':
        swap(r, 'Pi 5를 CU 트레이의 출력 보스 4개(Ø2.2 탭 구멍)에 바로 조이는 나사(4개 + 예비 2)',
             'Pi 5를 출력 받침 기둥 4개(Ø6 × 6, 구멍 Ø2.2 × 5 — ' + L2 + ')에 바로 조이는 나사(4개 + 예비 2)')
    if c == 'L58':
        swap(r, 'CU 트레이 보스에 M3×10(L34) 4개로 고정(R28)', '5 mm 출력 받침 기둥에 M3×5(L76) 4개로 고정(' + L2 + ', 6 mm 이상은 기둥 바닥을 뚫음)')
    if c == 'L10':
        r['role'] = ('스피커 파트 2개와 가운데 유닛의 판 16조각을 이 한 장에서 재단한다(' + L2 + ', 앞판 40°: 스피커 외형 273×128.5×z5~144.65, '
                     '안 부피 2.16 L). 스피커 옆판 128.5×139.65 ×4(40° 사선·앞 턱·통로 홈 27×22는 직접 톱으로), 뒤판 250×139.65 ×2, '
                     '윗판 250×14.14 ×2, 아랫판 250×90 ×2, 가운데 끝벽 117×56.35 ×2(ㄴ자 홈), 가운데 뒤판 702×56.35, 아랫판 679×90, '
                     '뚜껑 L·R 241×128.5 ×2. 배치 1181.8×385(톱날 3 mm 포함). 출처: body_L2.json plywood_cut_list')
        r['note'] = ('무료재단·무료배송. ' + L2 + ': 가게에 직사각형 16조각 배치를 맡기고 사선·홈·창은 직접 자른다. '
                     '가게가 이 배치를 못 하면 600×1200(이전 L63, 22,520원)을 다시 산다')
    if c == 'L07':
        r['role'] = ('좌우 스피커 유닛 2개(R25: 서브 없음). 스피커 파트(오꾸메 밀폐 상자, 안 부피 2.16 L · 유닛 뺀 2.06 L — ' + L2 + ')의 '
                     '수평에서 40° 기운 출력 앞판에 앞에서 단다(EVA 가스켓 3T + M4×12(L70) 4개 → 앞판 M4 열압입 인서트(L69) PCD115). '
                     '축이 가운데 귀와 17.3° 어긋남(W1 한계 20°). 앰프(L53)가 채널당 약 21 W(20 V)를 준다')
    if c == 'SP6':
        r['role'] = ('스피커 상자 2개(옆판·뒤판·윗판·아랫판, 안 부피 2.16 L ×2)와 가운데 유닛(아랫판·끝벽 2·뒤판) 접합, 출력 통로 틀·앞판 둘레 밀봉 '
                     '(앞판 아래면은 MS 폴리머). ' + L2 + ': 합판 11장 → 스피커 10 + 가운데 4. 합판은 본드로만 붙이고 마스킹테이프로 눌러 굳힌다(드릴 없음)')
    if c == 'L14':
        r['role'] = ('스피커 파트 2개(각 4개)와 가운데 유닛(6개)의 고무발 14개 + 예비 6. ' + L2 + ': 뒷바를 5 mm 띄우고 발은 모두 통로 뒤(y≥248)에 둔다; '
                     '모듈 USB 선은 뒷바 밑 닫힌 통로(y212~241 × z0~27)로 지나간다. 본체와 스피커 진동을 떼어 놓는다(D15)')
    if c == 'L42':
        r['role'] = (L2 + ': 방진 브래킷 BRK2-L/R(세운 핀 Ø6 z14~24)을 끝 부속 볼 뒷면 M3 자리 2곳씩에 고정하는 그로밋 4개 '
                     '(M3×16(L75) + M4 와셔(L36), 스피커를 올리기 전에 조임). 스피커는 핀에 얹히고 무게는 바닥 고무발이 받친다. 16개 남음')
    if c == 'L65':
        r['role'] = (L2 + ': 가운데 유닛 뚜껑 L·R(오꾸메) 2장을 붙여 두는 자석, 뚜껑마다 3쌍(뒤판 윗모서리 2 · 끝벽 윗모서리 1) = 12개 + 예비 2. '
                     '가운데 화면 뚜껑은 M3×10(L34) 4개로 고정하고 자석을 쓰지 않는다(D14). 극이 서로 당기게 순간접착제(ME8)로 붙인다')
    if c == 'L12':
        r['role'] = ('밀폐 스피커 상자 2개의 흡음재. ' + L2 + '(W1 조건): 상자마다 약 35 g을 상자 전체에 느슨하게, 바깥 끝 100 mm에 모으고 '
                     '유닛 뒤 10 mm·폴 벤트 축 10 mm는 비운다. 300 g이라 충분')
    if c == 'L17':
        r['role'] = ('모듈 7개 + PED의 USB-MIDI를 Pi 5 USB 한 포트로 모은다. 전원은 본체 5.1 V 강압 모듈(L58)에서 DC 잭으로 받는다(허브에 딸린 5 V 3 A 어댑터 선을 잘라 플러그 쪽을 쓴다). '
                     + L2 + ': 가운데 유닛 F칸 바닥에 뒤판에 붙여 놓고 포트는 앞(−y), 위에 앰프 선반. 딸린 0.6 m 업스트림 선(곧은 플러그)은 뚜껑 아래(z61.35)에 들어가지 않아 짧은 선(C25)을 쓴다')
        r['note'] = ('다나와에서 최저가 판매처로 이동해 구매. ' + L2 + ' 확인: 제품 사진상 업스트림은 USB 3.0 B 포트이고 DC 잭이 같은 끝면에 있다 '
                     '(CAD v16: 업스트림 B 포트와 DC 잭을 둘 다 −x 끝면(Pi·강압 쪽)에 그림, B y319 · DC y293 · z28.5, 5 V 선 약 225 — 포트 자리는 사진 추정). '
                     'DC 플러그(CAD v16): ㄱ자가 꼭 필요하지는 않다 — 딸린 곧은 플러그가 잭 면에서 32 mm 이내(몰딩 + 부트)이고 바로 뒤에서 R5쯤 위로 꺾을 수 있으면 쓴다. '
                     '32 mm보다 길면 IH190 젠더에 닿으니, 허브 잭 규격(5.5×2.1 / 5.5×2.5 / 3.5×1.35)과 배럴 길이를 재서 같은 규격 ㄱ자 DC 플러그 선을 산다(그때 값 추가)')
    if c == 'C04':
        r['note'] = (r['note'] + '. ' + L2 + ': 1 m ×8을 그대로 쓴다(남는 선 약 2.3 m를 앞 공간·통로에 감음). CAD가 권한 짧은 세트(1 m ×2 O1·O2, '
                     '0.5 m ×4 O3·O4·O7·PED, 0.3 m ×2 O5·O6, 같은 계열 개당 약 1,300~1,900원 — 사양 추정, 미확인)는 선택이며 사기 전에 끈으로 길이를 잰다')
    if c == 'C11':
        r['role'] = r['role'] + ' + ' + L2 + '·터치스크린 3판: 화면 받침 안 리본 누름 패드 자리(xr26.75~39.45, 사양 C-7)'
    if c == 'C10':
        r['role'] = r['role'] + ' + 터치스크린 3판: Waveshare MX1.25→3핀 전원선 이음이 경첩 고리 안에 오면 감쌈(사양 G-4)'
    if c == 'T01':
        r['role'] = ('제어 기판·센서 기판 손 납땜 + ' + L2 + ': 열압입 인서트 12개(스피커 앞판 M4 8 · 화면 뚜껑 레일 M3 4, 예비 따로)를 '
                     '기본 팁으로 넣는다(인서트 윗면을 판 면과 같게, 더 밀어 넣지 않음)')
    if c == 'CT2+':
        r['note'] = (r['note'] + '. ' + L2 + ': 뒷바 출력물 약 1.16 kg(L1 출력물을 빼서 대략 같음, 사양 추정) + 터치스크린 3판 출력물 약 98 g이 '
                     '이 스풀 안에 든다고 본다. 모자라면 1 kg 스풀 16,150원 추가')
    if c == 'C22':  # 보조배터리 선: 곧은 0.6 m → ㄱ자. 10/2: 주파집(개별 배송비) → Coms NC888 (사용자 요청: 엘레파츠 배송비 빼기, CAD 확인)
        r.update(name='[Coms] [NC888] Coms C타입 MM 꺾임 케이블 1m C to C', unit=2600, qty=1,
                 url='https://www.eleparts.co.kr/goods/view?no=18842596', seller='엘레파츠', section='엘레파츠', no_img=True,
                 role=('보조배터리 → 피아노 USB-C 입력(L57). ' + L2 + ': 보조배터리 +x 끝 USB-C에 꺾인 머리를 위(+z)로 꽂고(USB-C는 뒤집어 꽂을 수 있음), '
                       '선은 z55~60으로 +y를 따라 뒤판 케이블 통과(x405 z38)로 내려가 밖에서 PD 입력에 꽂는다. 남는 약 0.65 m는 배터리 위(z51~60.8)에 납작하게 감는다(CAD 10/2 확인). '
                       '조건: 보조배터리 USB-C 포트 중심이 바닥에서 약 15.5 mm 이상(아니면 CAD가 배터리 받침 +x 끝벽을 그 자리만 따냄)'),
                 note=('10/2 교체: 주파집 ㄱ자 0.5 m(4,270원)는 업체 직배송이라 개별 배송비 3,000원이 따로 붙고 엘레파츠 무료배송 합계에도 들어가지 않아 '
                       'Coms NC888(2,363.64원 VAT 별도 → 2,600원, 일반 배송, 평균 발송 3.5일, 1인 1개, C–C 양끝 90° 꺾임)로 바꿈 −4,670원. '
                       'C–C 선은 규격상 3 A라 20 V × 3 A = 60 W. 페이지에 W 정격·꺾임 면 표기가 없어 받으면 꺾임 방향과 20 V가 나오는지 확인'))
# ---- 10/1 L2 고침 (검토 반영): C24 ㄱ자 Pi 전원선, L1에만 있던 I/O 판·CU 트레이 글, SP6·L73 글
for r in ROWS:
    c = r['code']
    if c == 'C24':  # CAD README: 'Pi 5 전원선 MT666과 허브 DC 플러그는 ㄱ자형으로 그림. 곧은 플러그면 뚜껑 아래(z61.35)에 들어가지 않음'
        r.update(name='[Coms] [BT663] Coms USB 3.1 Type C 케이블 20cm USB 2.0 A to C타입 측면꺾임 (A 쪽을 잘라 씀)', unit=2550, qty=1,
                 url='https://www.eleparts.co.kr/goods/view?no=8105815', seller='엘레파츠', section='엘레파츠', no_img=True,
                 role=('5.1 V 강압 모듈(L58) 출력 → Pi 5 USB-C 전원 입력. ' + L2 + ': CAD는 Pi −y 가장자리에 측면 꺾임 USB-C 머리 14×13×9 '
                       '(Z-PI-PWR, x546~560 y249~262 z22~31)를 그리고 선을 −x로 뺐다 — 곧은 플러그는 뚜껑 아래(z61.35)에 들어가지 않는다(CAD README). '
                       'C 쪽이 옆으로 꺾인 20 cm 선의 A 쪽을 잘라 빨강(+)·검정(−)을 강압 출력 단자에 물리고 흰·초록 데이터선은 잘라 절연한다. '
                       'USB-C는 앞뒤가 같아 뒤집어 꽂으면 선이 +x로 나온다'),
                 note=(L2 + ' 교체(이전 MAXTEK MT666 DIY 2선, 1,190원 — 상품 사진상 곧은 플러그). 10/1 페이지: 2,318.18원(VAT 별도) → 2,550원, '
                       '평균 발송 3.5일(10/08), 1인 1개, 해외재고(개별 배송비 없음). 전원선 굵기는 공개되지 않음 — 20 cm라 3 A에서 약 0.1~0.25 V 떨어진다(추정). '
                       'Pi 저전압 경고가 뜨면 강압 출력(L58)을 조금 올린다. CAD v16: 머리는 사진 추정 14×12×7로 자리 14×13×9 안, 선은 −x(강압 쪽)로 꽂음 — −x 쪽이 비어 머리가 18까지 길어도 들어간다'))
    if c == 'L57':
        swap(r, 'I/O 판 안쪽 포켓에 끼워 USB-C 구멍이 밖을 향한다(R28)',
             L2 + ': 뒤판 출력물(PR-BACKPLATE)의 HUSB238 주머니(x415.8~426.2 z35.6~40.4)에 끼워 USB-C 입구(13 × 7.5, 바깥면에서 2 안)가 밖을 향한다(R28)')
    if c == 'L59':
        swap(r, 'I/O 판의 출력 구멍에 끼운다(스냅인, R28)',
             L2 + ': 뒤판 출력물의 2 mm 판 구멍(13.2 × 19.2)에 스냅인으로 끼우고 테두리는 2 mm 오목 자리에 들어가 바깥면과 같은 높이(R28)')
    if c == 'L15':
        swap(r, '② I/O 판의 댐퍼 페달 잭(R24) ③ CU 트레이의 앰프 입력 잭:',
             '② 뒤판 출력물의 댐퍼 페달 잭 J501(Ø6 구멍; J501 기판은 출력물 받침 위에 MS 폴리머, R24) '
             '③ 앰프 입력 잭 J702(J702 기판을 5 mm 받침 기둥에 M3×5(L76) 2개, 화면 뚜껑 아래 — ' + L2 + '):')
    if c == 'L53':
        swap(r, '입력은 CU 트레이의 앰프 입력 잭(L15 ③)에서 받는다',
             '입력은 J702 기판의 앰프 입력 잭(L15 ③)에서 받는다. ' + L2 + ': 허브 위 선반 받침에 M3×6(L35) 4개(받침 + 선반 관통)')
    if c == 'C14':
        swap(r, '② 헤드폰 잭 노멀 접점 → CU 트레이 앰프 입력 잭(L15 ③). 두 선 모두 CU 쪽 플러그를 남기고 끝 부속 쪽을 잘라 잭에 납땜한다(선이 끝 부속에 붙어 다니고 CU에서 뽑는다, R26)',
             '② 헤드폰 잭 노멀 접점 → J702 앰프 입력 잭(L15 ③, ' + L2 + ' 가운데 유닛). 두 선 모두 가운데 유닛 쪽 플러그를 남기고 끝 부속 쪽을 잘라 잭에 납땜한다'
             '(선이 끝 부속에 붙어 다니고 뒷바를 들 때 동글·J702에서 뽑는다, R26)')
    if c == 'L62':
        swap(r, '피아노 I/O 판의 USB-C 입력(L57)', '피아노 뒤판 출력물의 USB-C 입력(L57, ' + L2 + ')')
    if c == 'SP6':
        swap(r, '합판은 본드로만 붙이고 마스킹테이프로 눌러 굳힌다(드릴 없음)',
             '합판끼리는 나사 없이 본드로만 붙이고 마스킹테이프로 눌러 굳힌다. 판 구멍·홈·사선 가공은 따로 한다(Ø4 줄 구멍 홈 38개, Ø6·Ø8·Ø14 구멍, '
             '목재 인서트 막힌 구멍, 자석 자리 Ø8 × 3.2 12곳, 40° 사선·27 × 22 홈·뒤판 창 — 공구는 메모·선택기 v4wood)')
ROWS += [
    dict(row=None, section='엘레파츠', name='[ShenzenAV] 황동 인서트너트 M4x4x6 (내경 M4 · 높이 4 · 외경 6, 열압입)', unit=88, qty=10,
         role=(L2 + ': 스피커 출력 앞판(40°)의 유닛 나사 자리 PCD115에 4개씩 × 2 = 8 + 예비 2. CAD 구멍 Ø5.6 × 6.5(M4×6, 바깥 Ø5.6으로 그림)에 '
               '바깥 6을 열압입(물림 0.4, 보통 범위). CAD v16은 Ø6 × 4로 그림. M4×12(L70)가 플랜지 4 + 가스켓 3을 지나 5 들어가 인서트보다 1 더, 바닥까지 1.5(가스켓이 약 2로 눌리면 약 0.5). M4×12 이하만 — M4×14면 밀폐면을 뚫는다. 인서트는 면까지만'),
         url='https://www.eleparts.co.kr/goods/view?no=11045753', b='base', seller='엘레파츠', code='L69',
         note='10/1 페이지: 개당 80원(VAT 별도) → 88원, M.O.Q 1, 24시간 이내 발송. CAD 크기(M4×6, 바깥 Ø5.6)와 길이·바깥지름이 조금 다르다 — 구멍은 그대로 둔다'),
    dict(row=None, section='11번가', name='스텐 유두 렌치볼트 M4 × 12mm', unit=80, qty=10,
         role=(L2 + ': 스피커 유닛 CW-100B25를 앞판 인서트(L69)에 고정(유닛 플랜지 4 + EVA 가스켓 3을 지나 인서트에 물림), 4 × 2 = 8 + 예비 2. '
               'M4×12보다 긴 나사 금지(아래 구멍 밑 밀폐면 2.78 mm). 육각 렌치 2.5 mm(T11)'),
         url=GN_BOLT, b='base', seller='굿나잇몰', code='L70',
         note='10/1 옵션 값: 유두 M4 12mm 80원. 같은 상품의 옵션, 굿나잇몰 묶음배송'),
    dict(row=None, section='엘레파츠', name='[ShenzenAV] 황동 인서트너트 M3x4x4.5 (내경 M3 · 높이 4 · 외경 4.5, 열압입)', unit=88, qty=6,
         role=(L2 + '·터치스크린 3판: 화면 뚜껑 이음 레일 2개(SEAM-INS-1~4)의 윗면 y246·y322에 4개 + 예비 2. CAD 구멍 Ø4.0 × 4.5(바깥 Ø4로 그림)에 '
               '바깥 4.5를 열압입(물림 0.5, 보통 범위), 길이 4 = CAD 값(v16: Ø4.5 × 4로 그림) → M3×10(L34) 물림 4. 레일을 붙이기 전에 인두로 넣는다'),
         url='https://www.eleparts.co.kr/goods/view?no=11045752', b='base', seller='엘레파츠', code='L71',
         note='10/1 페이지: 개당 80원(VAT 별도) → 88원, M.O.Q 1, 24시간 이내 발송. 인서트 길이가 4가 아니면 자리파기 = 11.5 − (10 − 물림)으로 다시(사양 A-7)'),
    dict(row=None, section='나사코리아 nasakorea.com', name='평로렛 손잡이볼트 M4×20 스텐304 (5개, 머리 Ø8 × 6)', unit=6700, qty=1,
         role=(L2 + ': 스피커 ↔ 가운데 유닛 결합 — 가운데 안에서 끝벽의 고무 슬리브(L74)를 지나 스피커 안쪽 옆판의 M4 목재 인서트(L73)로, 한쪽 2개 × 2 = 4 + 예비 1. '
               '머리는 슬리브 끝에만 닿는다(D15). CAD v16: M4×20 = 인서트 물림 4, 최대 M4×22(물림 6, 끝이 바닥보다 2 위) — 더 길면 밀폐 판을 뚫는다. 손으로 돌림(공구 없이 스피커를 뗌)'),
         url='https://nasakorea.com/goods/goods_view.php?goodsNo=6090', b='base', seller='나사코리아 nasakorea.com', code='L72',
         note='10/1 페이지: 판매가 100원 + 옵션 "4*20/5개/스텐304" 6,600원 = 6,700원. 머리 Ø8 × 6(판매 사진 표) — CAD 머리 자리 Ø10 × 6.5 안. 밸런스 핀(v4pins) 주문과 같이 와서 배송비가 더 들지 않는다'),
    dict(row=None, section='나비엠알오', name='[Norelem] 셀프 태핑 나사산 인서트 07653-04 (스틸, 컷팅 보어) M4 · 겉지름 6.5 · 길이 6 (K94534159)', unit=1749, qty=5,
         role=(L2 + ': 스피커 안쪽 옆판(오꾸메 11.5)의 가운데 쪽 면에 2개씩 × 2 = 4 + 예비 1, 손잡이볼트(L72)가 물림. CAD v16 구멍 Ø6.0 × 8(깊이 멈춤, 밑 판 3.5). '
               '이 인서트는 길이 6 · 최소 구멍 깊이 8(제조사)이라 8 깊이 구멍에 맞고, 컷팅 보어형이라 슬롯형(07652)처럼 나사를 잠그지 않아 손잡이볼트를 손으로 풀 수 있다. '
               'M4×20 물림 4, M4×22 물림 6(= 인서트 길이). 구멍은 오꾸메 자투리에 Ø5.8부터 시험하고 헐거우면 Ø5.5(나무는 눌림)'),
         url='https://www.navimro.com/g/3544290/', b='base', seller='나비엠알오', code='L73',
         note=('10/1 CAD v16: 목재 퀵서트 전장 10(이전 L73, 572원)은 구멍 깊이 10이면 드릴 끝이 밀폐 판을 뚫어 쓸 수 없음 → 8 이하로. '
               '10/1 페이지: 1,590원(VAT 별도) → 1,749원, 1EA, 10/26 출하 예정, 반품 불가 주문품. 나비엠알오 공급가 55,690원으로 무료배송 유지. '
               '대안: 이전 L73(전장 10)의 슬롯 없는 위쪽 끝을 쇠톱으로 2 mm 잘라 씀(추가 0원) · 쿠팡 티알인서트 3020-040 M4×8 슬롯 260원(10/6 도착, 배송비 3,300) — 슬롯형은 최소 깊이 10이라 8 깊이 구멍에서 끝까지 안 들어갈 수 있음')),
    # 10/1 L2 고침: 엘레파츠 [NAVI] 3,210원 → 같은 상품을 나비엠알오에서 직접(2,959원) — 나비엠알오 주문이 공급가 5만 원을 넘어 배송비 3,300원도 빠진다
    dict(row=None, section='나비엠알오', name='[나비엠알오] 진공 실리콘 튜브 내경 4 · 외경 8 · 1 m 투명 (SL.STSM4, K15538680)', unit=2959, qty=1,
         role=(L2 + ': 고무 슬리브 — 13 mm로 4개 잘라 가운데 유닛 끝벽의 Ø8 구멍에 끼우고 손잡이볼트(L72)가 지나간다(쇠가 나무·출력물에 닿지 않게, D15). '
               'CAD의 "고무 호스 OD8 ID4.2" 대신. 남는 약 0.95 m는 예비'),
         url='https://www.navimro.com/g/312592/?k=K15538680', b='base', seller='나비엠알오', code='L74',
         note=('10/1 페이지: 2,690원(VAT 별도) → 2,959원, 1 m, 국산, 10/02 출하. 10/1 L2 고침: 엘레파츠 [NAVI] 같은 상품 3,210원에서 옮김 — '
               '나비엠알오 주문이 공급가 47,650 → 50,340원이 되어 기본 배송비 3,300원(5만 원 미만에만)이 빠진다')),
    dict(row=None, section='11번가', name='스텐 유두 렌치볼트 M3 × 16mm', unit=90, qty=6,
         role=(L2 + ': 방진 브래킷 BRK2-L/R을 끝 부속 볼 뒷면 M3 자리(볼 속 너트 L31)에 그로밋(L42) + M4 와셔(L36)로 고정, 2 × 2 = 4 + 예비 2. '
               'M3×10은 볼 속 너트에 닿지 않는다(L1부터 빠져 있던 줄). 육각 렌치 2 mm(T04), 스피커를 올리기 전에 조임'),
         url=GN_BOLT, b='base', seller='굿나잇몰', code='L75', note='10/1 옵션 값: 유두 M3 16mm 90원. 같은 상품의 옵션, 굿나잇몰 묶음배송'),
    dict(row=None, section='11번가', name='스텐 유두 렌치볼트 M3 × 5mm', unit=40, qty=10,
         role=(L2 + ': 가운데 유닛 보드를 5 mm 출력 받침 기둥(구멍 Ø2.5~2.7 × 4, 바닥 1.0)에 직접 탭 — 강압 모듈 4 · J702 잭 기판 2 · PED 보드 2 = 8 + 예비 2. '
               '6 mm 이상은 기둥 바닥을 뚫고 붙인 기둥을 밀어 올린다(L1의 CU 트레이 M3×10 대신)'),
         url=GN_BOLT, b='base', seller='굿나잇몰', code='L76', note='10/1 옵션 값: 유두 M3 5mm 40원. 같은 상품의 옵션, 굿나잇몰 묶음배송'),
    dict(row=None, section='11번가', name='스텐 접시머리 직결피스 8호(Ø4.2) 13mm', unit=40, qty=8,
         role=(L2 + ': 뒤판 출력물(PR-BACKPLATE) 바닥 탭 2 + 보조배터리 받침 4 = 6 + 예비 2. 구멍의 접시 자리는 모델에 있다(L49는 둥근머리)'),
         url='https://www.11st.co.kr/products/4350230243', b='base', seller='굿나잇몰', code='L77',
         note='10/1 옵션 값: 4.2mm×13mm 40원(같은 판매자의 접시머리 직결피스 상품). 굿나잇몰 묶음배송(주문 화면에서 묶음 여부 확인)'),
    dict(row=None, section='11번가', name='스텐 트러스 태핑 피스 1종 M3 × 12mm', unit=50, qty=4,
         role=(L2 + ': 허브 위 앰프 선반의 받침 2개를 가운데 뒤판(오꾸메)에 고정(받침 10 × 8에 8호가 들어가지 않음), 2 + 예비 2. '
               'CAD는 "Ø3 × 12~13 둥근머리 목재 나사, 머리 Ø6 이하" — 트러스 머리 약 Ø6이라 받으면 머리 자리를 확인'),
         url='https://www.11st.co.kr/products/4236198586', b='base', seller='굿나잇몰', code='L78',
         note='10/1 옵션 값: M3 12mm 50원(판매가 40 + 10). 굿나잇몰 묶음배송'),
    dict(row=None, section='엘레파츠', name='[코티니] 컴잇 벨크로 링벨트 케이블 선정리 [2X30cm] 블랙', unit=863, qty=1,
         role=(L2 + ': 보조배터리(선택, 105×71×32)를 출력 받침에 묶는 20 mm 벨크로 — 받침 앞·뒤 벽 바닥의 끈 구멍 22 × 3을 지나 배터리 밑·위를 감는다. '
               '보조배터리를 쓰지 않으면 빼도 된다'),
         url='https://www.eleparts.co.kr/goods/view?no=12816289', b='base', seller='엘레파츠', code='L79',
         note='10/1 페이지: 784.55원(VAT 별도) → 863원, 평균 발송 2.5일'),
    dict(row=None, section='엘레파츠', name='[Coms] [NA977] Coms USB Type A to Type B 변환 케이블 25cm USB 2.0 A 상향꺾임', unit=1530, qty=1,
         role=(L2 + ': 허브 업스트림 짧은 선(W209 대신) — Pi 5 위 USB-A ↔ 허브 업스트림. 딸린 0.6 m 곧은 선은 두 플러그 사이 21 mm · 뚜껑 아래 z50 안에 들어가지 않는다. '
               '허브 업스트림은 USB 3.0 B 포트(제품 사진) — USB 2.0 B 플러그도 꽂히고 USB-MIDI 8개에는 USB 2.0으로 충분. '
               'CAD v16 확정: Pi 5 GPIO 쪽 USB-A 2단의 아래 포트에 꽂는다(위 포트면 머리 윗면 z57로 뚜껑과 4.35뿐). A가 위로 올라가 R7로 꺾여 허브 선반 위 z56.5에서 U자로 돌고 x650.5에서 내려가 곧은 B로(길이 41 + 169 + 40 = 250). 남는 Pi USB-A는 RJ45 쪽 위 포트 하나. 받으면 꽂아서 A가 위로 꺾이는지 확인(아래로 꺾이는 선은 바닥에 닿아 못 씀)'),
         url='https://www.eleparts.co.kr/goods/view?no=4314354', b='consumable', seller='엘레파츠', code='C25',
         note='10/1 페이지: 1,390.91원(VAT 별도) → 1,530원, 평균 발송 3.5일. 엘레파츠에 "곧은 A + 꺾인 B" 짧은 선이 없어 고름(USB 3.0 A/B 양쪽 측면 꺾임 30 cm IH235는 6,120원)'),
]
# 10/1 L2 고침: 나비엠알오 배송비 규칙 — 기본 3,300원(VAT 포함), 공급가 5만 원(VAT 별도) 이상 무료(10/1 상품 페이지 배송 안내).
# 지금까지 선택기는 'O:나비엠알오' 묶음에 3,300원을 기준 없이 붙였다.
SHIP['O:나비엠알오'] = dict(fee=3300, b='base', thr=50000, vat=True, label='나비엠알오 (공급가 5만원 미만)')
ROWS_V4 =[r for r in ROWS if not r.get('drop')]
DROPPED = [r for r in ROWS if r.get('drop')]

# ---------------------------------------------------------------- categories and cards for existing rows
CATS = OrderedDict([
    ('key', ('건반 액션 (v4 신규)', 'v4 W1+ 4차(짧은 시소 건반 + 무게 레버 + 토션 스프링)에 새로 들어가는 부품. 3차에서 쓰던 레일·손나사·접시 스프링·탭 띠·납 등은 빠졌다(맨 아래 목록).')),
    ('cpu', ('제어 컴퓨터·전원', 'Pi 5, 저장장치, USB 허브, USB-C PD 전원부.')),
    ('sen', ('센서·MCU·배선', '건반 모듈 7개와 페달 보드의 감지 전자부와 배선.')),
    ('aud', ('오디오·스피커·뒷바', 'USB 동글 → TPA3110 앰프 → 스피커 2개, 스피커 상자와 가운데 유닛.')),
    ('mec', ('체결·펠트·접착', '나사, 펠트·EVA, 텅스텐 퍼티, 접착제.')),
    ('ped', ('페달·소프트웨어', '댐퍼 페달 1개와 무료 소프트웨어.')),
    ('fil', ('필라멘트', 'Bambu PETG Basic. v4에서 출력량이 늘었다.')),
    ('tool', ('공구', '제작에만 쓰고 피아노에는 남지 않는다. 기본 구성(100만원 한도) 밖.')),
])
# card id -> (cat, title, [codes], [old option ids])
CARDS_OLD = OrderedDict([
    ('pi', ('cpu', 'Raspberry Pi 5 2GB', ['L01'], ['CP1'])),
    ('heatsink', ('cpu', 'Pi 5 방열판', ['L02'], [])),
    ('piscrew', ('cpu', 'Pi 고정 나사 M2.5×6', ['L47'], [])),
    ('sd', ('cpu', 'microSD 카드', ['L03'], ['CPV1'])),
    ('hub', ('cpu', '전원형 USB 허브', ['L17'], ['CP8'])),
    ('charger', ('cpu', '65 W PD 벽 충전기', ['L62'], [])),
    ('pd', ('cpu', 'PD 트리거 (USB-C 20 V 입력)', ['L57'], [])),
    ('buck', ('cpu', '5.1 V 강압 모듈', ['L58'], [])),
    ('switch', ('cpu', '전원 스위치', ['L59'], [])),
    ('fuse', ('cpu', '퓨즈 홀더 + 퓨즈', ['L60'], [])),
    ('pdcable', ('cpu', 'USB-C PD 케이블 (보조배터리·벽 충전기)', ['C22', 'C23'], [])),
    ('pipower', ('cpu', 'Pi 전원선 (USB-C 제작용)', ['C24'], [])),
    ('mcu', ('sen', 'RP2040-Zero 8개 · 4067 먹스 7개', ['SEV2'], [])),
    ('hall', ('sen', '선형 홀센서 100개', ['L20'], ['SE1'])),
    ('magnet', ('sen', '네오디뮴 자석 Ø5×2 100개', ['L21'], ['SE3'])),
    ('sensorpcb', ('sen', '센서 기판 (만능기판 12×18)', ['L22'], [])),
    ('ctrlpcb', ('sen', '제어 기판 (만능기판 5×7)', ['SE8'], [])),
    ('header', ('sen', '핀헤더 1×40 (먹스·제로)', ['SE7'], [])),
    ('c100n', ('sen', '100nF 1206 콘덴서', ['SE4'], ['SEV4'])),
    ('c10u', ('sen', '10 µF 1206 콘덴서 (전원 안정)', ['L67'], [])),
    ('r1k1206', ('sen', '1 kΩ 1206 칩 저항 (끝 부속)', ['L68'], [])),
    ('extlead', ('sen', '끝 부속 6핀 잠금 커넥터 선', ['L66'], [])),
    ('usbcable', ('sen', '모듈 USB 케이블 8가닥', ['C04'], [])),
    ('ribbon', ('sen', '리본 케이블 (센서 → 제어 기판, 배선)', ['C05', 'SE12'], [])),
    ('shrink', ('sen', '열수축튜브', ['C10'], ['CT7'])),
    ('kapton', ('sen', '캡톤 테이프', ['C11'], [])),
    ('ties', ('sen', '케이블 타이', ['C12'], [])),
    ('solder', ('sen', '실납 0.8 mm 50 g', ['CT5'], [])),
    ('speaker', ('aud', '스피커 유닛 2개 (삼미 CW-100B25)', ['L07'], [])),
    ('amp', ('aud', 'TPA3110 앰프 보드', ['L53'], [])),
    ('ampcap', ('aud', '앰프 입력 필터 1 µF', ['L56'], [])),
    ('resist', ('aud', '저항 4종 (필터·풀다운)', ['L29'], [])),
    ('jack', ('aud', '3.5 mm 스테레오 잭', ['L15'], [])),
    ('dongle', ('aud', 'USB 오디오 동글 · 젠더', ['L51', 'L52'], ['AE7'])),
    ('xt30', ('aud', 'XT30 커넥터 수·암', ['L61'], [])),
    ('spkwire', ('aud', '스피커선 16AWG 7 m', ['C13'], ['AE8'])),
    ('aux', ('aud', '3.5 mm AUX 케이블', ['C14'], [])),
    ('fill', ('aud', '스피커 흡음 솜', ['L12'], ['SP5'])),
    ('feet', ('aud', '고무발 · 방진 그로밋', ['L14', 'L42'], ['SP4'])),
    ('ply400', ('aud', '오꾸메 합판 400×1200 (스피커 상자)', ['L10'], [])),
    ('ply600', ('aud', '오꾸메 합판 600×1200 (가운데 유닛)', ['L63'], [])),
    ('latch', ('aud', '토글 래치 2개', ['L64'], [])),
    ('lidmag', ('aud', '뚜껑 자석 Ø8×3', ['L65'], [])),
    ('woodscrew', ('aud', '직결피스 4종 (뒷바·스피커)', ['L49'], [])),
    ('woodglue', ('aud', '목공본드', ['SP6'], [])),
    ('nut', ('mec', 'M3 육각너트 (캡스턴용)', ['L31'], [])),
    ('m3x10', ('mec', 'M3×10 렌치볼트 (센서 바·CU)', ['L34'], [])),
    ('m3x6', ('mec', 'M3×6 렌치볼트 (제어 기판)', ['L35'], [])),
    ('washer', ('mec', 'M4 와셔', ['L36'], [])),
    ('felt3', ('mec', '양모 펠트 3T (앞 멈춤)', ['L39'], ['ME13', 'ME16'])),
    ('floor', ('mec', 'EVA 3T 바닥 시트 · 0.5T 접착 펠트', ['MEV1'], [])),
    ('tungsten', ('mec', '텅스텐 퍼티 (건반별 무게 조정)', ['ME14'], [])),
    ('ca', ('mec', '순간접착제', ['ME8'], [])),
    ('tape', ('mec', '종이 양면테이프', ['ME9'], [])),
    ('pedal', ('ped', '댐퍼 페달 (홀센서 개조)', ['L43'], [])),
    ('pedalcable', ('ped', '페달 신호선 3 m', ['L50'], [])),
    ('software', ('ped', 'FluidSynth · Salamander 음원 (무료)', ['L45', 'L46'], [])),
    ('filament', ('fil', 'PETG 필라멘트 8 kg (흰 2 · 검정 6)', ['CT2', 'CT2+'], ['CT3'])),
    ('iron', ('tool', '납땜 인두 세트 (Pinecil)', ['T01', 'T02', 'T03'], ['CTV3', 'CT17', 'CT10'])),
    ('meter', ('tool', '멀티미터', ['CT11'], [])),
    ('stripper', ('tool', '정밀 와이어 스트리퍼', ['CT21'], [])),
    ('nipper', ('tool', '니퍼', ['CT13'], [])),
    ('caliper', ('tool', '버니어 캘리퍼스', ['CT12'], [])),
    ('hexkey', ('tool', '육각 렌치 2.5 mm', ['T11'], [])),
])
# ---- 10/1 L2: new cards (placed after a related card) and retitled cards
def _insert_after(key, new):
    items = list(CARDS_OLD.items())
    i = [k for k, _ in items].index(key) + 1
    CARDS_OLD.clear()
    CARDS_OLD.update(items[:i] + new + items[i:])


_insert_after('pipower', [('pbstrap', ('cpu', '보조배터리 고정 벨크로 (L2)', ['L79'], [])),
                          ('hubup', ('cpu', '허브 업스트림 짧은 선 (L2)', ['C25'], []))])
_insert_after('speaker', [('spkmount', ('aud', '스피커 유닛 고정 — M4 열압입 인서트 + M4×12 (L2)', ['L69', 'L70'], []))])
_insert_after('feet', [('podjoint', ('aud', '스피커 ↔ 가운데 결합 — 손잡이볼트·목재 인서트·실리콘 슬리브 (L2)', ['L72', 'L73', 'L74'], []))])
_insert_after('m3x10', [('lidins', ('mec', 'M3 열압입 인서트 (화면 뚜껑 레일, L2)', ['L71'], [])),
                        ('brkbolt', ('mec', 'M3×16 렌치볼트 (BRK2 브래킷, L2)', ['L75'], [])),
                        ('m3x5', ('mec', 'M3×5 렌치볼트 (가운데 보드 받침, L2)', ['L76'], []))])
CARDS_OLD['woodscrew'] = ('aud', '직결피스·태핑 피스 (고무발·그릴·받침, L2)', ['L49', 'L77', 'L78'], [])
CARDS_OLD['m3x10'] = ('mec', 'M3×10 렌치볼트 (센서 바·화면 뚜껑)', ['L34'], [])
CARDS_OLD['feet'] = ('aud', '고무발 · 방진 그로밋 (BRK2)', ['L14', 'L42'], [])
CARDS_OLD['pipower'] = ('cpu', 'Pi 전원선 (측면 꺾임 USB-C, A 쪽을 잘라 씀 — L2)', ['C24'], [])  # 10/1 L2 고침

# why "사용 안함" matters (short)
UNUSED = {
    'pi': '음원 컴퓨터가 없어진다. 가진 Pi 5나 PC·노트북을 음원 호스트로 쓸 때만 고른다.',
    'sd': '가진 microSD(8 GB 이상)를 쓸 때.',
    'hub': '가진 전원형 USB 허브(8포트 이상)를 쓸 때. 없으면 모듈을 Pi에 다 꽂을 수 없다.',
    'charger': '65 W 이상 USB-C PD 충전기가 있으면(노트북 충전기 등).',
    'mcu': '빼면 건반 모듈이 동작하지 않는다. 가진 RP2040-Zero가 있을 때만.',
    'hall': '빼면 건반을 감지할 수 없다. 가진 선형 홀센서가 있을 때만.',
    'magnet': '빼면 센서가 건반을 볼 수 없다. 가진 Ø5×2 자석이 있을 때만.',
    'speaker': '내장 스피커 없이 헤드폰만 쓴다(R13·R14 요구를 포기). 앰프·스피커 상자도 함께 빼는 것이 맞다.',
    'amp': '내장 스피커를 쓰지 않을 때만.',
    'ply400': '스피커 상자를 만들지 않을 때만(내장 스피커 포기).',
    'ply600': '가운데 유닛(건반 보관함·CU·보조배터리 칸)을 만들지 않을 때만. Pi·앰프를 둘 자리를 따로 마련해야 한다.',
    'filament': '가진 PETG가 충분할 때만(흰 약 1.6 kg, 검정 약 5.4 kg 필요).',
    'pedal': '댐퍼 페달 없이 연주한다(R24 포기) 또는 가진 서스테인 페달을 쓸 때.',
    'tungsten': '건반별 무게 편차가 작으면 필요 없다. 단계 0 시편 뒤 정해도 된다.',
    'felt3': '빼면 건반 바닥이 딱딱하게 부딪혀 소리가 난다. 가진 3T 펠트가 있을 때만.',
    'c10u': '빼면 센서 기판·끝 부속의 전원 안정 콘덴서(△1·△3)를 달 수 없다. 가진 1206 10 µF(16 V 이상)가 16개 있을 때만.',
    'r1k1206': '빼면 끝 부속 R401~R403·R411을 달 수 없다. 가진 1206 1 kΩ가 4개 있을 때만.',
    # 10/1 L2
    'spkmount': '빼면 스피커 유닛을 앞판에 달 수 없다(L2에서 직결피스 16 mm가 빠짐). 가진 M4 열압입 인서트(바깥 Ø5.6~6)와 M4×12가 8개씩 있을 때만.',
    'podjoint': '빼면 스피커와 가운데 유닛을 이을 수 없다(도브테일·토글 래치가 없어짐). 가진 M4 손잡이 나사·목재 인서트·고무 슬리브가 4개씩 있을 때만.',
    'lidins': '빼면 화면 뚜껑을 레일에 나사로 고정할 수 없다. 가진 M3 열압입 인서트(길이 4)가 4개 있을 때만.',
    'brkbolt': '빼면 BRK2 방진 브래킷을 볼에 고정할 수 없다(M3×10은 볼 속 너트에 닿지 않음). 가진 M3×16 버튼헤드가 4개 있을 때만.',
    'm3x5': '빼면 강압 모듈·J702·PED 보드를 받침 기둥에 달 수 없다. 가진 M3×5가 8개 있을 때만(6 mm 이상은 기둥 바닥을 뚫음).',
    'pbstrap': '보조배터리를 쓰지 않거나 가진 20 mm 벨크로가 있을 때.',
    'hubup': '빼면 허브를 Pi에 이을 선이 없다(딸린 0.6 m 선은 뚜껑 아래에 들어가지 않음). 가진 짧은 USB A–B 꺾임 선이 있을 때만.',
}
CARD_ROLE = {  # card summary when the first row's text is not enough
    'resist': '1 kΩ: R701·R702·R501·R551 / 100 kΩ: R301(O1·O7)·R302·R303·R502 / 100 Ω: R503(100 Ω ×5 = 20 Ω, 필수) / 10 kΩ: 회로에서 안 씀(예비)',
}
TOOL_UNUSED = '가지고 있거나 빌릴 수 있으면 고른다. 피아노에는 남지 않는 공구다.'

# overrides for old option copy that describes v3-only situations
V4NOTE = {
    'CT3': 'v4에서는 PETG가 스프링 역할을 하지 않아(건반은 강철 봉 위 시소, 복귀는 강철 추), 필라멘트 물성 편차와 크리프가 동작에 주는 영향이 v3보다 훨씬 작다. 대신 치수 정밀도(노치·가이드 폭)는 여전히 중요하다.',
    'ME13': 'v4에서 3T 펠트는 건반 앞 멈춤과 키퍼 훅 밑에 쓴다. 훅은 평소 닿지 않는 키퍼라 v3보다 부담이 작지만, 앞 멈춤 펠트의 반발(e)은 단계 0에서 확인한다.',
    'ME16': 'v4의 앞 멈춤은 흑건 노치 들림과 유령 재타건에 영향을 준다(앞 펠트 반발 e ≤ 0.22가 단계 0 합격선). EVA는 반발이 커서 v4에서는 더 권하지 않는다.',
}


def row_item(r):
    key = None
    if r.get('no_img'):  # 10/1 L2: replaced product without a photo — keep the old line's photo off it
        pass
    elif r.get('img_path'):
        key = img_file('P-' + r['code'], r['img_path'])
    elif re.match(r'^[LCT]\d+$', r['code']):
        key = img_line(r['code'])
    else:
        o = OPT.get(r['code'].rstrip('+'))
        if o:
            for i, it in enumerate(o['add']):
                if it['name'][:14] == str(r['name'])[:14]:
                    key = img_opt(r['code'].rstrip('+'), i)
                    break
            if key is None and o['add']:
                key = img_opt(r['code'].rstrip('+'), 0)
    return dict(name=r['name'], qty=r['qty'], unit=r['unit'], b=r['b'], store=r['seller'] or r['section'],
                ship=ship_key(r['seller'], r['section']), url=r['url'], role=r['role'], note=r['note'], img=key,
                code=r['code'], verified=True)


def opt_items(oid):
    o = OPT[oid]
    out = []
    for i, it in enumerate(o['add']):
        out.append(dict(name=it['name'], qty=it['qty'], unit=it['unit'], b=it['b'], store=it['store'],
                        ship=ship_key(it['store']) or (('O:' + it['store']) if it['ship'] else None),
                        ship_fee=it['ship'], url=it['url'], role='', note=it.get('note') or '', img=img_opt(oid, i),
                        code=oid, verified=bool(it['price_ok'] and it['page_ok'])))
    return out


def old_option(oid, card_rows):
    o = OPT[oid]
    d = DRAFT.get(oid, {})
    items = opt_items(oid)
    # options that replace only some lines of a card keep the other current rows
    keep = [row_item(r) for r in card_rows if r['code'] not in o['replaces']]
    if oid == 'CT3':  # scale 5 kg set to v4 8 kg (white 2, black 6)
        for it in items:
            if '블랙' in it['name']:
                it['qty'] = 6
    return dict(id=oid, title=d.get('tagline') or o['title'], items=keep + items, pros=d.get('pros', []), cons=d.get('cons', []),
                note=V4NOTE.get(oid, ''), verified=all(i['verified'] for i in items), src='v3 절감안')


CARDS = []
_PRIOR_CARDS = {c['id']: c for c in _PRIOR['cards']} if CACHED else {}
for cid, (cat, title, codes, opts) in CARDS_OLD.items():
    rows = [r for r in ROWS_V4 if r['code'] in codes]
    if not rows:
        continue
    cur = dict(id='cur', title='현재 목록', items=[row_item(r) for r in rows], pros=[], cons=[], note='', verified=True, src='')
    if CACHED and cid in _PRIOR_CARDS:
        pc = _PRIOR_CARDS[cid]
        pimg = {(it['code'], str(it['name'])): it.get('img') for it in pc['options'][0]['items']}
        for it in cur['items']:
            if not it['img']:
                it['img'] = pimg.get((it['code'], str(it['name'])))
        rolemap = {(r['code'], str(r['name'])): r['role'] for r in rows}
        olds = []
        for o in pc['options']:
            if o.get('src') == 'v3 절감안':
                o = json.loads(json.dumps(o))
                by_code = {}
                for r_ in rows:
                    by_code.setdefault(r_['code'], []).append(r_)
                for j, it in enumerate(o['items']):
                    k = (it['code'], str(it['name']))
                    if k in rolemap and it.get('role'):
                        it['role'] = rolemap[k]
                    elif it.get('role') and len(by_code.get(it['code'], [])) == 1:
                        o['items'][j] = row_item(by_code[it['code']][0])  # the one source row was changed (e.g. 9/30 vendor)
                olds.append(o)
        options = [cur] + olds
    else:
        options = [cur] + [old_option(o, rows) for o in opts if o in OPT]
    CARDS.append(dict(id=cid, cat=cat, title=title, role=CARD_ROLE.get(cid, rows[0]['role']), options=options,
                      unused=UNUSED.get(cid, TOOL_UNUSED if cat == 'tool' else '이미 가지고 있거나 다른 방법으로 해결할 때 고른다.')))

for c in CARDS:  # 9/30: SP6 moved to 나비엠알오 — share the option-level shipping group used by the v4 cards
    for o in c['options']:
        for it in o['items']:
            if it['code'] == 'SP6' and '나비엠알오' in str(it.get('store', '')):
                it['ship'], it['ship_fee'] = 'O:나비엠알오', 3300
            # 10/1 L2: new rows at sellers whose shipping lives in the v4 cards' option-level groups
            if it['code'] in ('L73', 'L74'):  # L74: 10/1 L2 고침으로 나비엠알오로 옮김
                it['ship'], it['ship_fee'] = 'O:나비엠알오', 3300
            # 10/1 L2 고침: 주파집 ㄱ자 선은 해외재고 "개별 배송비 3,000원" — 엘레파츠 무료배송 기준과 따로 붙는다
            if it['code'] == 'C22' and '주파집' in str(it['name']):
                it['ship'], it['ship_fee'] = 'O:엘레파츠 주파집(개별배송)', 3000
            if it['code'] == 'L72':
                it['ship'], it['ship_fee'] = 'O:나사코리아 nasakorea.com', 3300

# ---- circuit review (2026-09-28): picker-only alternatives (the xlsx keeps the current line)
CABLE = os.path.join(V4, 'cable')
CAB = {c['name'].split(' — ')[0]: c for c in json.load(open(os.path.join(CABLE, 'result.json')))['candidates']}


def alt_item(name, qty, unit, b, store, ship, url, note, img, code, fee=0):
    return dict(name=name, qty=qty, unit=unit, b=b, store=store, ship=ship, ship_fee=fee, url=url, role='', note=note, img=img,
                code=code, verified=True)


cv, nm = CAB['CVILUX DH-20M50052'], CAB['NETmate NM-GCM01BN']
EXTRA_OPTS = {
    'usbcable': [
        dict(id='dh20m', title='CVILUX DH-20M50052 8개 — 몰드 치수가 공개된 유일한 선 (DigiKey)', src='케이블 조사', verified=True, fits=True,
             items=[alt_item('CVILUX DH-20M50052 USB 2.0 A(수) → C(수) 1 m 검정 (DigiKey 2987-DH-20M50052-ND)', 8, cv['price_krw'], 'consumable',
                             'DigiKey Korea (digikey.kr)', 'DigiKey', cv['url'],
                             f'C 몰드 {cv["plug_w_mm"]} × {cv["plug_t_mm"]} × {cv["plug_l_mm"]}(DigiKey 캘리퍼 실측 사진) — 한계 12.5 × 7.6 × 25 안',
                             img_file('P-C04-dh20m', cv['image_file']), 'C04', fee=cv['shipping_krw'])],
             pros=['C 쪽 몰드 폭·두께·길이가 모두 공개 실측값이고 한계 안에 든다', '홈 가장자리·건반 꼬리와의 틈이 지금보다 늘어난다'],
             cons=['DigiKey 배송비 20,000원(6만원 미만). 8개 + 배송비 40,744원으로 NA993보다 27,784원 많다'],
             note='10개는 22,009원(개당 2,200.9원)이라 1,265원 더 내면 예비 2개가 생긴다. DigiKey에서 다른 부품과 합쳐 6만원을 넘기면 배송비가 없다. 받으면 단계 0 시험 19에 그대로 쓴다.'),
        dict(id='gcm01', title='NETmate NM-GCM01BN 8개 — 길이 약 32로 맞지 않음', src='케이블 조사', verified=True, fits=False,
             items=[alt_item('NETmate NM-GCM01BN USB-A 2.0 → C 1 m 블랙', 8, nm['price_krw'], 'consumable', '엘레파츠', '엘레파츠', nm['url'],
                             '폭 11.5·두께 6은 공개 도면 값. 몰드 길이는 공개되지 않았고 도면 축척으로 잰 추정 약 32(부트 포함)',
                             img_file('P-C04-gcm01', nm['image_file']), 'C04')],
             pros=['엘레파츠 합배송이라 배송비가 더 붙지 않는다'],
             cons=['추정 길이 약 32가 한계 25를 넘는다. 인터페이스를 "딱딱한 몰드 ≤ 25, 부트 포함 ≤ 35"로 고칠 때만 쓸 수 있다'],
             note=''),
    ],
    'c10u': [
        dict(id='cl31k16', title='삼성 CL31A106KOHNNNE 20개 (같은 16 V, LCSC 해외 재고)', src='9/28 조사', verified=True, fits=True,
             items=[alt_item('Samsung CL31A106KOHNNNE (1206 10 µF 16 V X5R)', 20, 242, 'base', '아이씨뱅큐 (LCSC 해외 재고, P014969585)', '아이씨뱅큐',
                             'https://www.icbanq.com/P014969585', '10~99개 220원 + VAT = 242원, 1개 단위. 16개 + 예비 4',
                             img_file('P-L67-cl31', os.path.join(CIRC, 'img_10uF_tme.jpg')), 'L67')],
             pros=['같은 규격(16 V X5R)을 필요한 만큼만 산다', '제조사·품번이 분명하다'],
             cons=['해외 재고라 1주 이내 도착, 취소·반품 불가'], note=''),
        dict(id='cl31k25', title='삼성 CL31A106KAHNNNE 29개 (25 V, Arrow 해외 재고)', src='9/28 조사', verified=True, fits=True,
             items=[alt_item('Samsung CL31A106KAHNNNE (1206 10 µF 25 V X5R)', 29, 44, 'base', '아이씨뱅큐 (Arrow 해외 재고, P011854568)', '아이씨뱅큐',
                             'https://www.icbanq.com/P011854568', '29개 단위, 40원 + VAT = 44원. 9/28 재고 29개뿐',
                             img_file('P-L67-cl31', os.path.join(CIRC, 'img_10uF_tme.jpg')), 'L67')],
             pros=['가장 싸다', '25 V라 전압 여유가 더 크다'],
             cons=['도착이 1달 이내로 느리다', '재고가 29개뿐이라 주문할 때 없을 수 있다'], note=''),
    ],
}
# ---- 10/1 L2 고침: 선택기 전용 대안 (구매 목록 xlsx는 현재 줄을 그대로 둔다)
_c23 = next(r for r in ROWS_V4 if r['code'] == 'C23')
_c24img = next((it.get('img') for it in _PRIOR_CARDS.get('pipower', {}).get('options', [{}])[0].get('items', []) if it['code'] == 'C24'), None) if CACHED else None
EXTRA_OPTS['pdcable'] = [
    dict(id='jupazip', title='보조배터리 선을 주파집 ㄱ자 C–C 0.5 m로 (개별 배송비 3,000원)', src='10/1 L2', verified=True, fits=True,
         items=[alt_item('[주파집] C to C 90도 ㄱ자 60W 고속충전 케이블 싱글 [0.5m]', 1, 4270, 'consumable', '엘레파츠', 'O:엘레파츠 주파집(개별배송)',
                         'https://www.eleparts.co.kr/goods/view?no=14097908',
                         '10/1 페이지: 3,881.82원(VAT 별도) → 4,270원, 한쪽 90° 꺾임, 60 W 표기. 업체 직배송이라 개별 배송비 3,000원이 따로 붙고 엘레파츠 무료배송 합계에 들어가지 않는다', None, 'C22', fee=3000),
                row_item(_c23)],
         pros=['60 W 표기, 0.5 m라 남는 선이 적다', '한쪽만 꺾여 PD 입력 쪽은 곧은 플러그'],
         cons=['개별 배송비 3,000원', 'NC888보다 4,670원 비싸다'], note=''),
]
EXTRA_OPTS['pipower'] = [
    dict(id='mt666', title='이전: MAXTEK MT666 곧은 DIY 2선 (CAD 자리에 맞지 않음)', src='10/1 이전 목록', verified=True, fits=False,
         fitlabel='곧은 플러그 — CAD 확인 필요',
         items=[alt_item('[MAXTEK] 맥스텍 USB C타입 제작용 케이블 2선 Type-C Male(숫) DIY PD 전원 [MT666]', 1, 1190, 'consumable', '엘레파츠', '엘레파츠',
                         'https://www.eleparts.co.kr/goods/view?no=17358965', '10/1 페이지: 1,081.82원(VAT 별도) → 1,190원. 상품 사진상 곧은 USB-C 플러그',
                         _c24img, 'C24')],
         pros=['1,360원 싸고 두 선이 이미 벗겨져 있다'],
         cons=['곧은 플러그라 CAD의 측면 꺾임 머리 14×13×9 자리(Z-PI-PWR)에 맞지 않는다 — CAD README: 곧은 플러그는 뚜껑 아래 z61.35에 들어가지 않음',
               'CAD가 곧은 플러그 자리를 따로 낼 때만 고른다'],
         note='검토에서 권한 Coms IH627 "꺾임 젠더 C to C"는 상세 사진상 수–수 U자(180°) 어댑터라 MT666과 Pi 사이에 끼울 수 없어 넣지 않았다.'),
]
for c in CARDS:
    c['options'] += EXTRA_OPTS.get(c['id'], [])

# ---- v4 small patches on existing cards
_PV = os.path.join(V4, 'parts', 'parts.json')  # 10/2: scratch v4/parts was purged by the OS tmp cleaner — fall back to the known T04 values
PARTS_V = ({it['id']: it for it in json.load(open(_PV))['items']} if os.path.exists(_PV)
           else {'T04': {'url': 'https://www.11st.co.kr/products/4236198413', 'image_file': None}})
for c in CARDS:
    for o in c['options']:
        for it in o['items']:
            if it['code'] == 'L31' and not it['img']:
                it['img'] = img_file('P-m3nut', os.path.join(V4, 'parts', 'img', 'm3_nut.jpg'))
    if c['id'] == 'hexkey':
        t = PARTS_V['T04']
        c['title'] = '육각 렌치 2.5 mm · 2 mm'
        c['options'][0]['items'].append(dict(name='육각 렌치 2 mm (볼트 상품 옵션)', qty=1, unit=600, b='tool', store='굿나잇몰', ship='굿나잇몰',
            url=t.get('url', ''), role='v4 캡스턴 M3 버튼헤드 조정용', note='v4에서 추가 (세트스크루용 1.5 mm 대신)', img=img_file('P-hex', t.get('image_file')), code='T04', verified=True))

json.dump(dict(cards=CARDS, dropped=[dict(code=r['code'], name=r['name'], reason=r['drop'], price=r['unit'] * r['qty']) for r in DROPPED],
               ship=SHIP, cats=CATS), open(os.path.join(HERE, 'data_old.json'), 'w'), ensure_ascii=False, indent=1)
json.dump(IMGS, open(os.path.join(HERE, 'imgs_old.json'), 'w'))
print('cards', len(CARDS), 'imgs', len(IMGS), 'dropped', [d['code'] for d in DROPPED], 'CACHED' if CACHED else 'live')
